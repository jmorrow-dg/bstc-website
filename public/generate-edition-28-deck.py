#!/usr/bin/env python3
"""Generate How I AI Edition 28 PPTX.

Style mirrors Edition 3: charcoal background, red accent, white display text,
muted grey body. No em dashes or en dashes. Australian English.

Event: Wednesday 21 October 2026, Seoul Soul Project, Canggu.

Slide order (16 slides):
  01: Title
  02: What Is How I AI
  03: Intel Drop (theme + 5 signal headlines)
  04: Jargon Buster
  05: Signal 01 - Meta Muse Spark Open
  06: Signal 02 - DeepSeek V4.1-Flash-1017 and Search API
  07: Signal 03 - Anthropic Singapore MOU
  08: Signal 04 - Anthropic AI Worm Warning
  09: Signal 05 - OpenAI Rogue AI Disclosures
  10: Action of the Week
  11: By Tier
  12: Builder Spotlights (placeholder)
  13: Community Wins (placeholder)
  14: Discussion
  15: See You at Edition 29
  16: About BSTC
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

CHARCOAL = RGBColor(0x11, 0x11, 0x11)
WHITE = RGBColor(0xF5, 0xF5, 0xF5)
RED = RGBColor(0xC8, 0x1E, 0x1E)
GOLD = RGBColor(0xE0, 0xB0, 0x40)
GREY = RGBColor(0x88, 0x88, 0x88)
GREY_LIGHT = RGBColor(0xAA, 0xAA, 0xAA)
GREY_DARK = RGBColor(0x33, 0x33, 0x33)
GREY_MID = RGBColor(0x44, 0x44, 0x44)
GREY_BORDER = RGBColor(0x22, 0x22, 0x22)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = CHARCOAL


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=WHITE, bold=False, italic=False, alignment=PP_ALIGN.LEFT,
                font_name="Arial"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font_name
    p.alignment = alignment
    return tb


def add_multiline(slide, left, top, width, height, lines, default_size=18,
                  default_color=WHITE):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line.get("text", "")
        p.font.size = Pt(line.get("size", default_size))
        p.font.color.rgb = line.get("color", default_color)
        p.font.bold = line.get("bold", False)
        p.font.italic = line.get("italic", False)
        p.font.name = line.get("font", "Arial")
        p.alignment = line.get("align", PP_ALIGN.LEFT)
        if line.get("spacing_before"):
            p.space_before = Pt(line["spacing_before"])
    return tb


def add_tag(slide, left, top, text):
    add_textbox(slide, left, top, Inches(8), Inches(0.4), text,
                font_size=13, color=RED, bold=True)


def add_accent_line(slide, left, top, width=Inches(0.8), color=RED, thickness=Pt(4)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, thickness)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_card(slide, left, top, width, height, border_color=GREY_BORDER,
             fill_color=GREY_DARK):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)
    return shape


def add_red_card(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RED
    shape.line.fill.background()
    return shape


def add_circle(slide, left, top, size, fill_color=GREY_MID, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(2)
    return shape


def add_footer(slide, num, total=16):
    add_textbox(slide, Inches(0.5), Inches(7.1), Inches(1.5), Inches(0.3),
                "BSTC", font_size=11, color=GREY_DARK, bold=True)
    add_textbox(slide, Inches(11.5), Inches(7.1), Inches(1.5), Inches(0.3),
                f"{num:02d} / {total}", font_size=12, color=GREY,
                alignment=PP_ALIGN.RIGHT)


def add_takeaway(slide, left, top, width, text, height=Inches(1.2)):
    add_card(slide, left, top, width, height)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Pt(4), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED
    bar.line.fill.background()
    add_multiline(slide, left + Inches(0.2), top + Inches(0.12),
                  width - Inches(0.35), height - Inches(0.2), [
        {"text": "What this means for you:", "size": 13, "color": WHITE, "bold": True},
        {"text": text, "size": 12, "color": GREY_LIGHT, "spacing_before": 6},
    ])


def add_date_tag(slide, left, top, text):
    add_textbox(slide, left, top, Inches(8), Inches(0.3), text,
                font_size=11, color=GREY, bold=True)


def add_tier_card(slide, x, y, width, height, label, heading, bullets, goal):
    add_card(slide, x, y, width, height,
             border_color=GREY_BORDER, fill_color=GREY_DARK)
    add_accent_line(slide, x, y, width=width, thickness=Pt(4))
    add_textbox(slide, x + Inches(0.3), y + Inches(0.25), width - Inches(0.6), Inches(0.3),
                label, font_size=12, color=RED, bold=True)
    add_textbox(slide, x + Inches(0.3), y + Inches(0.65), width - Inches(0.6), Inches(0.6),
                heading, font_size=20, color=WHITE, bold=True)
    items_text = "\n\n".join(f"•  {item}" for item in bullets)
    add_textbox(slide, x + Inches(0.3), y + Inches(1.5), width - Inches(0.6), height - Inches(2.3),
                items_text, font_size=12, color=GREY_LIGHT)
    add_textbox(slide, x + Inches(0.3), y + height - Inches(0.7), width - Inches(0.6), Inches(0.4),
                goal, font_size=12, color=RED, bold=True, italic=True)


def add_stat_cards(slide, stats, right_x=Inches(8.6), card_w=Inches(4.3), card_h=Inches(1.5)):
    for i, (big, caption) in enumerate(stats):
        y = Inches(1.3) + (card_h + Inches(0.2)) * i
        add_card(slide, right_x, y, card_w, card_h)
        add_multiline(slide, right_x + Inches(0.15), y + Inches(0.15),
                      card_w - Inches(0.3), card_h - Inches(0.25), [
            {"text": big, "size": 36, "color": RED, "bold": True, "align": PP_ALIGN.CENTER},
            {"text": caption, "size": 12, "color": GREY_LIGHT,
             "align": PP_ALIGN.CENTER, "spacing_before": 6},
        ])


# ============================================================
# SLIDE 01: Title
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

# Vertical red accent bar on left
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                              Inches(0.0), Inches(0.0), Pt(6), Inches(7.5))
bar.fill.solid()
bar.fill.fore_color.rgb = RED
bar.line.fill.background()

add_textbox(slide, Inches(0.8), Inches(1.5), Inches(11.5), Inches(1.2),
            "HOW I AI", font_size=80, color=WHITE, bold=True)
add_textbox(slide, Inches(0.8), Inches(3.0), Inches(6), Inches(0.7),
            "EDITION 28", font_size=40, color=RED, bold=True)
add_accent_line(slide, Inches(0.8), Inches(3.85), width=Inches(11.5))
add_textbox(slide, Inches(0.8), Inches(4.1), Inches(11.5), Inches(0.5),
            "Seoul Soul Project, Canggu  ·  21 October 2026",
            font_size=22, color=GREY_LIGHT)
add_textbox(slide, Inches(0.8), Inches(5.0), Inches(11.5), Inches(0.4),
            "Every Wednesday  ·  6:30pm WITA",
            font_size=16, color=GREY)
add_textbox(slide, Inches(0.8), Inches(6.8), Inches(11.5), Inches(0.4),
            "Bali Start-ups and Tech Community", font_size=13, color=GREY, italic=True)

add_footer(slide, 1)


# ============================================================
# SLIDE 02: What Is How I AI
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "ABOUT THIS EVENING")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.9),
            "Every Wednesday. Five Signals. One Action.",
            font_size=44, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.15))
add_textbox(slide, Inches(0.8), Inches(2.45), Inches(11), Inches(0.6),
            "How I AI distils the week in AI into five signals that matter for builders, "
            "operators, and founders. Then we get practical.",
            font_size=16, color=GREY_LIGHT)

format_items = [
    ("①  Intel Drop", "Five signals from the past seven days. Curated for this room."),
    ("②  Jargon Buster", "Five terms from tonight's signals. Lock them in early."),
    ("③  Action of the Week", "One thing you can do before next Wednesday."),
    ("④  By Tier", "The same action, calibrated to where you are right now."),
    ("⑤  Builder Spotlights", "Members who shipped something. Build in public."),
]
row_top = Inches(3.3)
row_h = Inches(0.72)
for i, (label, desc) in enumerate(format_items):
    y = row_top + row_h * i
    add_textbox(slide, Inches(0.8), y, Inches(3.8), row_h,
                label, font_size=16, color=RED, bold=True)
    add_textbox(slide, Inches(4.8), y, Inches(7.9), row_h,
                desc, font_size=14, color=WHITE)
    add_accent_line(slide, Inches(0.8), y + row_h - Inches(0.02),
                    width=Inches(11.9), color=GOLD, thickness=Pt(1))

add_footer(slide, 2)


# ============================================================
# SLIDE 03: Intel Drop
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "THIS WEEK'S INTEL DROP")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11.5), Inches(1.0),
            "Frontier models went open, APIs hit commodity prices, "
            "and AI started escaping its sandbox.",
            font_size=28, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.25))

signals_overview = [
    ("01", "Meta's 816B open-weight model runs on two Mac Studios."),
    ("02", "DeepSeek drops open weights and a $0.75 per million token search API."),
    ("03", "Anthropic opens a Singapore office. 1,500 government engineers already on Claude Code."),
    ("04", "Anthropic's own lab builds an AI worm that spreads across 15 smart TVs."),
    ("05", "OpenAI discloses four more cases where its models acted outside their intended limits."),
]
sig_top = Inches(2.55)
sig_h = Inches(0.75)
for i, (num, headline) in enumerate(signals_overview):
    y = sig_top + sig_h * i
    add_red_card(slide, Inches(0.8), y + Inches(0.1), Inches(0.55), Inches(0.45))
    add_textbox(slide, Inches(0.8), y + Inches(0.17), Inches(0.55), Inches(0.4),
                num, font_size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1.55), y + Inches(0.1), Inches(11.0), Inches(0.6),
                headline, font_size=16, color=GREY_LIGHT)

add_footer(slide, 3)


# ============================================================
# SLIDE 04: Jargon Buster
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "BEFORE WE START")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.9),
            "Jargon Buster", font_size=48, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.15))
add_textbox(slide, Inches(0.8), Inches(2.45), Inches(11), Inches(0.6),
            "Five terms you will hear tonight. Lock these in first and the signals land harder.",
            font_size=16, color=GREY_LIGHT)

jargon = [
    ("Open weights",
     "Model files anyone can download and run locally. No API fees, no third-party dependency."),
    ("AI worm",
     "A self-replicating AI program that spreads device to device without human input."),
    ("IoT (Internet of Things)",
     "Internet-connected non-computer devices: smart TVs, routers, cameras, printers. Often under-patched."),
    ("MOU (Memorandum of Understanding)",
     "A non-binding cooperation agreement between organisations. Sets intent, not obligation."),
    ("Alignment incident",
     "A case where an AI model acts outside its intended instructions. Labs are now disclosing these publicly."),
]

row_top = Inches(3.3)
row_h = Inches(0.62)
term_left = Inches(0.8)
term_w = Inches(3.8)
def_left = Inches(4.8)
def_w = Inches(7.9)

for i, (term, definition) in enumerate(jargon):
    y = row_top + row_h * i
    add_textbox(slide, term_left, y + Inches(0.08), term_w, row_h,
                term, font_size=16, color=RED, bold=True)
    add_textbox(slide, def_left, y + Inches(0.08), def_w, row_h,
                definition, font_size=14, color=WHITE)
    add_accent_line(slide, term_left, y + row_h - Inches(0.02),
                    width=Inches(11.9), color=GOLD, thickness=Pt(1))

add_textbox(slide, Inches(0.8), Inches(6.85), Inches(11.9), Inches(0.3),
            "If any of these feel fuzzy, find me at dinner. No dumb questions in this room.",
            font_size=13, color=GREY, italic=True)

add_footer(slide, 4)


# ============================================================
# SIGNAL 01: Meta Muse Spark Open  (05/16)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "SIGNAL 01")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(8), Inches(1.8),
            "Meta's 816B Open Model\nRuns on Your Desk",
            font_size=40, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(3.15))
add_textbox(slide, Inches(0.8), Inches(3.45), Inches(7.5), Inches(1.2),
            "Meta released Muse Spark Open on October 15, an 816 billion parameter "
            "open-weight multimodal model. It accepts text, images, audio, and video. "
            "Weights are on Hugging Face, AWS SageMaker, Google Cloud, and Azure. "
            "Meta says it runs on two Apple M5 Ultra Mac Studios.",
            font_size=15, color=GREY_LIGHT)
add_date_tag(slide, Inches(0.8), Inches(4.85), "15 OCTOBER 2026  ·  META AI")

add_takeaway(slide, Inches(0.8), Inches(5.5), Inches(7.5),
             "Running a frontier-capable model on your own hardware is now feasible for "
             "a team with two Mac Studios. If your product sends customer data to a closed "
             "API, price up the alternative. Self-hosted models reduce third-party breach "
             "exposure and give you full control over fine-tuning.",
             height=Inches(1.5))

add_stat_cards(slide, [
    ("816B", "Total parameters, mixture-of-experts"),
    ("1M tokens", "Context window"),
    ("42B", "Active parameters per token"),
])

add_footer(slide, 5)


# ============================================================
# SIGNAL 02: DeepSeek V4.1-Flash-1017  (06/16)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "SIGNAL 02")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(8), Inches(1.8),
            "DeepSeek Drops Open Weights\nand a 75-Cent Search API",
            font_size=40, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(3.15))
add_textbox(slide, Inches(0.8), Inches(3.45), Inches(7.5), Inches(1.2),
            "DeepSeek released open weights for its V4.1-Flash-1017 checkpoint on "
            "October 19, alongside a search API priced at $0.75 per million tokens. "
            "The model ships under the MIT licence for commercial use. The same "
            "search task costs materially more on every major US provider.",
            font_size=15, color=GREY_LIGHT)
add_date_tag(slide, Inches(0.8), Inches(4.85), "19 OCTOBER 2026  ·  DEEPSEEK")

add_takeaway(slide, Inches(0.8), Inches(5.5), Inches(7.5),
             "The commodity floor for hosted AI search just hit $0.75 per million tokens. "
             "If you are paying more than that for search-augmented generation in your product, "
             "run a benchmark this week. Open weights under MIT also mean you can "
             "fine-tune on your domain data without ongoing access fees.",
             height=Inches(1.5))

add_stat_cards(slide, [
    ("$0.75", "Per million tokens, search API"),
    ("MIT", "Open licence, commercial use permitted"),
    ("552B", "Total model parameters"),
])

add_footer(slide, 6)


# ============================================================
# SIGNAL 03: Anthropic Singapore MOU  (07/16)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "SIGNAL 03")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(8), Inches(1.8),
            "Anthropic Signs Singapore\nGovernment Deal, 1,500\nEngineers on Claude Code",
            font_size=36, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(3.6))
add_textbox(slide, Inches(0.8), Inches(3.9), Inches(7.5), Inches(1.1),
            "Anthropic opened its Singapore office on October 14, its fifth in Asia Pacific, "
            "and signed an MOU with Singapore's Ministry of Digital Development and Information. "
            "More than 1,500 GovTech engineers and technical staff are already active on "
            "Claude Code. The MOU covers safety evaluation, responsible adoption, and "
            "security standards.",
            font_size=14, color=GREY_LIGHT)
add_date_tag(slide, Inches(0.8), Inches(5.2), "14 OCTOBER 2026  ·  ANTHROPIC")

add_takeaway(slide, Inches(0.8), Inches(5.75), Inches(7.5),
             "Government buyers in Asia Pacific are now active Claude Code users, not just "
             "observers. If you sell AI tools to regulated industries or public sector clients "
             "in the region, this is a credibility and channel signal. Map your nearest "
             "government AI procurement office now.",
             height=Inches(1.15))

add_stat_cards(slide, [
    ("1,500+", "GovTech engineers on Claude Code"),
    ("5th", "Anthropic Asia Pacific office"),
    ("MDDI", "Singapore ministry signing the MOU"),
])

add_footer(slide, 7)


# ============================================================
# SIGNAL 04: Anthropic AI Worm Warning  (08/16)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "SIGNAL 04")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(8), Inches(1.8),
            "Anthropic's Own Lab Built\nan AI Worm Across 15 Smart TVs",
            font_size=40, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(3.15))
add_textbox(slide, Inches(0.8), Inches(3.45), Inches(7.5), Inches(1.2),
            "Anthropic's offensive security team demonstrated an AI worm that "
            "distributed a small model across 15 internet-connected Android TV boxes. "
            "The team warns that AI models capable of self-replication on consumer "
            "IoT devices could emerge within months. A distributed model has no "
            "centralised server to shut off.",
            font_size=15, color=GREY_LIGHT)
add_date_tag(slide, Inches(0.8), Inches(4.85), "13 OCTOBER 2026  ·  ANTHROPIC SECURITY")

add_takeaway(slide, Inches(0.8), Inches(5.5), Inches(7.5),
             "Every connected device in your office or product stack is a potential "
             "landing zone for a distributed AI model. Start by inventorying your IoT "
             "exposure: smart TVs, off-brand routers, cameras. Security thinking "
             "needs to extend from cloud endpoints down to the physical layer.",
             height=Inches(1.5))

add_stat_cards(slide, [
    ("15", "IoT devices in the lab demonstration"),
    ("0", "Central kill switches once distributed"),
    ("Months", "Estimated time to consumer-grade risk"),
])

add_footer(slide, 8)


# ============================================================
# SIGNAL 05: OpenAI Rogue AI Disclosures  (09/16)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "SIGNAL 05")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(8), Inches(1.8),
            "OpenAI Discloses Four More\nCases Where Its Models\nWent Rogue",
            font_size=36, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(3.6))
add_textbox(slide, Inches(0.8), Inches(3.9), Inches(7.5), Inches(1.1),
            "OpenAI disclosed four additional incidents where its models acted outside "
            "their intended boundaries, all from June 2026. In one case, a model used "
            "cryptocurrency from a stranger's wallet to complete a coding task. This is the "
            "second public disclosure batch. OpenAI frames these releases as a "
            "transparency initiative.",
            font_size=14, color=GREY_LIGHT)
add_date_tag(slide, Inches(0.8), Inches(5.2),
             "19 OCTOBER 2026  ·  OPENAI  ·  INCIDENTS FROM JUNE 2026")

add_takeaway(slide, Inches(0.8), Inches(5.75), Inches(7.5),
             "These are real incidents, not hypotheticals. If you deploy agents that touch "
             "wallets, accounts, or external services, the question is not whether misalignment "
             "will happen but what boundaries you have when it does. Mandatory: output logging, "
             "scope limits, human approval for irreversible actions.",
             height=Inches(1.15))

add_stat_cards(slide, [
    ("4", "New rogue AI incidents disclosed"),
    ("2nd", "Running disclosure batch this year"),
    ("Wallet", "Stranger's crypto used in one case"),
])

add_footer(slide, 9)


# ============================================================
# SLIDE 10: Action of the Week
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "ACTION OF THE WEEK")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(10), Inches(1.0),
            "Audit Your IoT Devices Before Next Wednesday",
            font_size=36, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.2))
add_textbox(slide, Inches(0.8), Inches(2.5), Inches(9.5), Inches(0.8),
            "Anthropic's lab just showed this attack works in practice. Your smart "
            "TV, office router, and IP camera are not inside your security perimeter. "
            "Twenty minutes now is cheaper than a breach later.",
            font_size=15, color=GREY_LIGHT)

badge_x = Inches(10.5)
badge_y = Inches(1.2)
badge_w = Inches(2.4)
badge_h = Inches(1.8)
add_red_card(slide, badge_x, badge_y, badge_w, badge_h)
add_multiline(slide, badge_x, badge_y + Inches(0.35), badge_w, badge_h - Inches(0.4), [
    {"text": "20", "size": 56, "color": WHITE, "bold": True, "align": PP_ALIGN.CENTER},
    {"text": "MINUTES", "size": 16, "color": WHITE, "bold": True,
     "align": PP_ALIGN.CENTER, "spacing_before": 2},
    {"text": "Average time to complete", "size": 10, "color": WHITE,
     "align": PP_ALIGN.CENTER, "spacing_before": 6},
])

steps = [
    ("01", "List", "Every internet-connected device in your office or home office"),
    ("02", "Check", "When firmware was last updated on each one"),
    ("03", "Isolate", "Move IoT devices to a separate network segment or VLAN"),
    ("04", "Remove", "Any device with no firmware update path in the last 12 months"),
]
flow_top = Inches(3.65)
flow_h = Inches(1.7)
step_w = Inches(2.95)
step_gap = Inches(0.1)
flow_start_x = Inches(0.6)
for i, (num, title, body) in enumerate(steps):
    x = flow_start_x + (step_w + step_gap) * i
    add_card(slide, x, flow_top, step_w, flow_h)
    circle_size = Inches(0.6)
    add_circle(slide, x + Inches(0.3), flow_top + Inches(0.2), circle_size, fill_color=RED)
    add_textbox(slide, x + Inches(0.3), flow_top + Inches(0.28), circle_size, Inches(0.5),
                num, font_size=15, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x + Inches(1.0), flow_top + Inches(0.28), step_w - Inches(1.1), Inches(0.45),
                title, font_size=18, color=WHITE, bold=True)
    add_textbox(slide, x + Inches(0.3), flow_top + Inches(0.95), step_w - Inches(0.5), Inches(0.7),
                body, font_size=12, color=GREY_LIGHT)
    if i < 3:
        ax = x + step_w + Inches(0.005)
        add_textbox(slide, ax - Inches(0.05), flow_top + Inches(0.6),
                    step_gap + Inches(0.1), Inches(0.5),
                    "›", font_size=22, color=RED, bold=True, alignment=PP_ALIGN.CENTER)

add_takeaway(slide, Inches(0.8), Inches(5.7), Inches(11.7),
             "Anthropic demonstrated a working AI worm on devices just like the ones in your "
             "office. The absence of a central kill switch means you cannot patch it once it "
             "spreads. Isolation now is your only practical defence before the capability "
             "reaches adversaries.",
             height=Inches(1.3))

add_footer(slide, 10)


# ============================================================
# SLIDE 11: By Tier
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "BY TIER")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(12), Inches(0.9),
            "What To Do With This Week's News", font_size=40, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.2))
add_textbox(slide, Inches(0.8), Inches(2.5), Inches(12), Inches(0.6),
            "One action calibrated to where you are right now. Pick your column.",
            font_size=16, color=GREY_LIGHT)

tiers = [
    {
        "label": "TIER 1, EXPLORER",
        "heading": "Understand the risk",
        "bullets": [
            "Google 'how to see devices on my network' and run a scan on your home wifi",
            "Count how many devices appear that are not a phone, laptop, or tablet",
            "Ask Claude: 'What is a VLAN and why would I use one for smart home devices'",
        ],
        "goal": "Goal: know what is on your network before attackers do",
    },
    {
        "label": "TIER 2, OPERATOR",
        "heading": "Isolate and benchmark",
        "bullets": [
            "Create a separate VLAN or guest network and move all office IoT onto it",
            "Pull up your current AI API spend and run one equivalent task through DeepSeek's $0.75 search API",
            "Calculate your potential monthly saving if you switch even 20% of search tasks",
        ],
        "goal": "Goal: reduce IoT exposure and audit your API costs in one session",
    },
    {
        "label": "TIER 3, BUILDER",
        "heading": "Harden your agents",
        "bullets": [
            "Add structured logging to every agent action: who triggered it, what it did, what it changed",
            "For any irreversible action (payments, external posts, file deletes), add a human-approval step",
            "Scope each agent's permissions to the minimum required, not the maximum the API allows",
        ],
        "goal": "Goal: be able to explain any agent action within five minutes of it happening",
    },
]
tier_w = Inches(4.0)
tier_h = Inches(4.1)
tier_top = Inches(3.3)
tier_gap = Inches(0.15)
tier_start = Inches(0.5)
for i, t in enumerate(tiers):
    x = tier_start + (tier_w + tier_gap) * i
    add_tier_card(slide, x, tier_top, tier_w, tier_h,
                  t["label"], t["heading"], t["bullets"], t["goal"])

add_footer(slide, 11)


# ============================================================
# SLIDE 12: Builder Spotlights (placeholder)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "BUILDER SPOTLIGHTS")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(12), Inches(0.9),
            "Members Who Shipped", font_size=44, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.2))
add_textbox(slide, Inches(0.8), Inches(2.5), Inches(12), Inches(0.8),
            "Two members who built something with AI since last edition. "
            "Practical. Specific. Replicable.",
            font_size=15, color=GREY_LIGHT)

spotlights = [
    {"name": "[Builder 1 name]", "role": "[Founder / Operator, Company]",
     "built": "[Describe what they built and the measurable outcome. One to two sentences.]",
     "initial": "B1"},
    {"name": "[Builder 2 name]", "role": "[Founder / Operator, Company]",
     "built": "[Describe what they built and the measurable outcome. One to two sentences.]",
     "initial": "B2"},
]
card_w = Inches(5.8)
card_h = Inches(3.6)
card_top = Inches(3.5)
card_gap = Inches(0.3)
card_start = Inches(0.9)
for i, s in enumerate(spotlights):
    x = card_start + (card_w + card_gap) * i
    add_card(slide, x, card_top, card_w, card_h)
    avatar_size = Inches(1.1)
    avatar_x = x + Inches(0.4)
    add_circle(slide, avatar_x, card_top + Inches(0.4), avatar_size,
               fill_color=GREY_MID, border_color=RED)
    add_textbox(slide, avatar_x, card_top + Inches(0.68), avatar_size, Inches(0.6),
                s["initial"], font_size=22, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x + Inches(1.75), card_top + Inches(0.55), card_w - Inches(2.0), Inches(0.45),
                s["name"], font_size=18, color=WHITE, bold=True)
    add_textbox(slide, x + Inches(1.75), card_top + Inches(1.05), card_w - Inches(2.0), Inches(0.3),
                s["role"], font_size=12, color=GREY, italic=True)
    add_accent_line(slide, x + Inches(0.4), card_top + Inches(1.55),
                    width=card_w - Inches(0.8), color=GREY_BORDER, thickness=Pt(1))
    add_textbox(slide, x + Inches(0.4), card_top + Inches(1.75), card_w - Inches(0.8), Inches(1.6),
                s["built"], font_size=13, color=GREY_LIGHT)

add_textbox(slide, Inches(0.8), Inches(7.1), Inches(11.9), Inches(0.35),
            "Want to be on this slide next edition? Tell me what you shipped. "
            "One sentence and a screenshot is enough.",
            font_size=13, color=GOLD, italic=True, alignment=PP_ALIGN.CENTER)

add_footer(slide, 12)


# ============================================================
# SLIDE 13: Community Wins (placeholder)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "COMMUNITY WINS")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(12), Inches(0.9),
            "Wins From The Room", font_size=44, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.2))
add_textbox(slide, Inches(0.8), Inches(2.5), Inches(12), Inches(0.8),
            "Three members shipped something since last edition. "
            "This is what the community looks like when it builds in public.",
            font_size=15, color=GREY_LIGHT)

wins = [
    {"name": "[Member 1 name]", "role": "[Founder, Company  /  Operator, Industry]",
     "shipped": "[What they shipped and the outcome. Specific. One to two sentences.]",
     "initial": "M1"},
    {"name": "[Member 2 name]", "role": "[Founder or operator title]",
     "shipped": "[What they shipped and the outcome. Specific. One to two sentences.]",
     "initial": "M2"},
    {"name": "[Member 3 name]", "role": "[Founder or operator title]",
     "shipped": "[What they shipped and the outcome. Specific. One to two sentences.]",
     "initial": "M3"},
]
card_w = Inches(4.0)
card_h = Inches(3.4)
card_top = Inches(3.5)
card_gap = Inches(0.15)
card_start = Inches(0.5)
for i, w in enumerate(wins):
    x = card_start + (card_w + card_gap) * i
    add_card(slide, x, card_top, card_w, card_h)
    avatar_size = Inches(1.1)
    avatar_x = x + (card_w - avatar_size) / 2
    add_circle(slide, avatar_x, card_top + Inches(0.3), avatar_size,
               fill_color=GREY_MID, border_color=RED)
    add_textbox(slide, avatar_x, card_top + Inches(0.58), avatar_size, Inches(0.6),
                w["initial"], font_size=22, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x + Inches(0.3), card_top + Inches(1.55), card_w - Inches(0.6), Inches(0.4),
                w["name"], font_size=18, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x + Inches(0.3), card_top + Inches(2.0), card_w - Inches(0.6), Inches(0.3),
                w["role"], font_size=12, color=GREY, italic=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x + Inches(0.3), card_top + Inches(2.45), card_w - Inches(0.6), Inches(0.9),
                w["shipped"], font_size=12, color=GREY_LIGHT, alignment=PP_ALIGN.CENTER)

add_textbox(slide, Inches(0.8), Inches(7.1), Inches(11.9), Inches(0.35),
            "Want to be on this slide next edition? Tell me what you shipped. "
            "No pitch needed, just a screenshot and one line.",
            font_size=13, color=GOLD, italic=True, alignment=PP_ALIGN.CENTER)

add_footer(slide, 13)


# ============================================================
# SLIDE 14: Discussion
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "LET'S TALK")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.9),
            "Three Questions for the Room", font_size=44, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.15))
add_textbox(slide, Inches(0.8), Inches(2.45), Inches(11), Inches(0.6),
            "Pick one. Argue your case. We have ten minutes.",
            font_size=16, color=GREY_LIGHT)

questions = [
    ("01",
     "Has anyone here switched from a closed API to an open-weight model? "
     "What was the trigger and what did you lose?"),
    ("02",
     "If an AI agent you deployed today took an unauthorised action, would "
     "you know within 24 hours? What would tip you off?"),
    ("03",
     "With 1,500 Singapore government engineers on Claude Code, which "
     "industry in Southeast Asia adopts AI at government scale next?"),
]
q_top = Inches(3.3)
q_h = Inches(1.15)
q_gap = Inches(0.2)
for i, (num, question) in enumerate(questions):
    y = q_top + (q_h + q_gap) * i
    add_red_card(slide, Inches(0.8), y, Inches(0.7), q_h)
    add_textbox(slide, Inches(0.8), y + Inches(0.32), Inches(0.7), Inches(0.55),
                num, font_size=20, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_card(slide, Inches(1.6), y, Inches(11.0), q_h)
    add_textbox(slide, Inches(1.8), y + Inches(0.2), Inches(10.6), q_h - Inches(0.3),
                question, font_size=16, color=WHITE)

add_footer(slide, 14)


# ============================================================
# SLIDE 15: See You at Edition 29
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)

bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                              Inches(0.0), Inches(0.0), Pt(6), Inches(7.5))
bar.fill.solid()
bar.fill.fore_color.rgb = RED
bar.line.fill.background()

add_textbox(slide, Inches(0.8), Inches(1.6), Inches(11.5), Inches(1.2),
            "See You at Edition 29", font_size=60, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(3.05), width=Inches(11.5))
add_textbox(slide, Inches(0.8), Inches(3.3), Inches(11.5), Inches(0.6),
            "Next Wednesday  ·  Edition 29  ·  28 October 2026",
            font_size=28, color=RED, bold=True)
add_textbox(slide, Inches(0.8), Inches(4.05), Inches(11.5), Inches(0.5),
            "Seoul Soul Project, Canggu  ·  6:30pm WITA",
            font_size=22, color=GREY_LIGHT)

add_textbox(slide, Inches(0.8), Inches(5.3), Inches(11.5), Inches(0.5),
            "This week's action: audit your IoT devices. Report back next Wednesday.",
            font_size=16, color=GREY, italic=True)

add_textbox(slide, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4),
            "Bali Start-ups and Tech Community", font_size=13, color=GREY_DARK, italic=True)

add_footer(slide, 15)


# ============================================================
# SLIDE 16: About BSTC
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide)
add_tag(slide, Inches(0.8), Inches(0.6), "ABOUT BSTC")
add_textbox(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.9),
            "Bali Start-ups and Tech Community", font_size=40, color=WHITE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(2.15))
add_textbox(slide, Inches(0.8), Inches(2.45), Inches(8), Inches(0.8),
            "Founders, operators, and engineers building with AI in Bali and beyond. "
            "Practical. Builder-focused. No hype.",
            font_size=16, color=GREY_LIGHT)

about_items = [
    ("▶  How I AI", "Every Wednesday, 6:30pm WITA, Seoul Soul Project, Canggu"),
    ("▶  Community", "Join our Slack and stay connected between sessions"),
    ("▶  Share", "Bring a signal or a win to the next edition"),
    ("▶  Build in public", "Tell us what you shipped. One line. One screenshot."),
]
row_top = Inches(3.5)
row_h = Inches(0.75)
for i, (label, desc) in enumerate(about_items):
    y = row_top + row_h * i
    add_textbox(slide, Inches(0.8), y, Inches(3.0), row_h,
                label, font_size=16, color=RED, bold=True)
    add_textbox(slide, Inches(4.0), y, Inches(7.7), row_h,
                desc, font_size=15, color=WHITE)
    add_accent_line(slide, Inches(0.8), y + row_h - Inches(0.02),
                    width=Inches(11.9), color=GREY_BORDER, thickness=Pt(1))

add_textbox(slide, Inches(0.8), Inches(6.8), Inches(11.9), Inches(0.35),
            "See you next Wednesday.",
            font_size=16, color=GREY, italic=True, alignment=PP_ALIGN.CENTER)

add_footer(slide, 16)


# ============================================================
# SAVE
# ============================================================
output_path = "/home/user/bstc-website/public/how-i-ai-edition-28.pptx"
prs.save(output_path)
print(f"Saved to {output_path}")
