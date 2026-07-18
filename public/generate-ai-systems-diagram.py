#!/usr/bin/env python3
"""Premium BSTC-brand diagram: AI Systems at the centre of the organisation.

Refined hub-and-spoke framework rendered as a hand-built SVG with Lucide-style
thin-line icons and Inter typography, then screenshotted by Chrome at 2x for
crisp output. Dark node discs with muted gold rings and soft glow (no loud
solid fills); red reserved as a single accent on Security & Governance. Matches
the How I AI deck: charcoal #111, gold + red accents, 13.333 x 7.5 (16:9).

Usage:
  python3 generate-ai-systems-diagram.py svg    # writes .svg + .html
  python3 generate-ai-systems-diagram.py pptx   # embeds .png into a .pptx slide
"""

import os, sys

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.join(HERE, "bstc-ai-systems-diagram")

GOLD   = "#C9A24A"   # muted, editorial gold (between BSTC #B8860B and deck #E0B040)
GOLD_HI = "#E0B040"
RED    = "#C03A2B"   # softened brand red, used once
ICON   = "#EAE3D2"   # warm off-white icon stroke
WARM   = "#F4F1E9"   # display white
TITLE  = "#CDB36A"   # node titles
BODY   = "#9A958C"   # descriptions
FAINT  = "#55504A"   # credits

# ---- Lucide icon path bodies (24x24) ----
ICONS = {
    "compass": '<circle cx="12" cy="12" r="10"/><polygon points="16.24 7.76 14.12 14.12 7.76 16.24 9.88 9.88 16.24 7.76"/>',
    "heart": '<path d="M19 14c1.49-1.46 3-3.21 3-5.5A5.5 5.5 0 0 0 16.5 3c-1.76 0-3 .5-4.5 2-1.5-1.5-2.74-2-4.5-2A5.5 5.5 0 0 0 2 8.5c0 2.3 1.5 4.05 3 5.5l7 7Z"/>',
    "trending": '<polyline points="22 7 13.5 15.5 8.5 10.5 2 17"/><polyline points="16 7 22 7 22 13"/>',
    "gear": '<path d="M12.22 2h-.44a2 2 0 0 0-2 2v.18a2 2 0 0 1-1 1.73l-.43.25a2 2 0 0 1-2 0l-.15-.08a2 2 0 0 0-2.73.73l-.22.38a2 2 0 0 0 .73 2.73l.15.1a2 2 0 0 1 1 1.72v.51a2 2 0 0 1-1 1.74l-.15.09a2 2 0 0 0-.73 2.73l.22.38a2 2 0 0 0 2.73.73l.15-.08a2 2 0 0 1 2 0l.43.25a2 2 0 0 1 1 1.73V20a2 2 0 0 0 2 2h.44a2 2 0 0 0 2-2v-.18a2 2 0 0 1 1-1.73l.43-.25a2 2 0 0 1 2 0l.15.08a2 2 0 0 0 2.73-.73l.22-.39a2 2 0 0 0-.73-2.73l-.15-.08a2 2 0 0 1-1-1.74v-.5a2 2 0 0 1 1-1.74l.15-.09a2 2 0 0 0 .73-2.73l-.22-.38a2 2 0 0 0-2.73-.73l-.15.08a2 2 0 0 1-2 0l-.43-.25a2 2 0 0 1-1-1.73V4a2 2 0 0 0-2-2z"/><circle cx="12" cy="12" r="3"/>',
}


def icon(name, cx, cy, box=40, color=ICON):
    s = box / 24.0
    return (f'<g transform="translate({cx - box/2:.2f},{cy - box/2:.2f}) scale({s:.4f})" '
            f'fill="none" stroke="{color}" stroke-width="1.05" '
            f'stroke-linecap="round" stroke-linejoin="round">{ICONS[name]}</g>')


def node(cx, cy, name):
    return (
        f'<circle cx="{cx}" cy="{cy}" r="52" fill="none" stroke="{GOLD}" '
        f'stroke-width="3" opacity="0.22" filter="url(#glow)"/>'
        f'<circle cx="{cx}" cy="{cy}" r="50" fill="url(#disc)" stroke="{GOLD}" stroke-width="1.4"/>'
        f'<circle cx="{cx}" cy="{cy}" r="50" fill="none" stroke="#000" stroke-opacity="0.25" stroke-width="0.6"/>'
        + icon(name, cx, cy)
    )


def tspans(lines, x, dy=20):
    out = []
    for i, ln in enumerate(lines):
        out.append(f'<tspan x="{x}" dy="{0 if i==0 else dy}">{ln}</tspan>')
    return "".join(out)


def build_svg():
    HUBX, HUBY, HUBR = 640, 372, 165
    nodes = {"lead": (560, 112), "ops": (560, 632), "del": (300, 372), "grow": (980, 372)}

    parts = []
    parts.append(f'''<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">
<defs>
  <radialGradient id="bg" cx="50%" cy="40%" r="80%">
    <stop offset="0%" stop-color="#191711"/><stop offset="55%" stop-color="#121212"/><stop offset="100%" stop-color="#0B0B0B"/>
  </radialGradient>
  <radialGradient id="disc" cx="50%" cy="38%" r="75%">
    <stop offset="0%" stop-color="#1F1F1F"/><stop offset="100%" stop-color="#151515"/>
  </radialGradient>
  <radialGradient id="hub" cx="50%" cy="36%" r="80%">
    <stop offset="0%" stop-color="#1C1B18"/><stop offset="100%" stop-color="#121212"/>
  </radialGradient>
  <linearGradient id="goldring" x1="0" y1="0" x2="0" y2="1">
    <stop offset="0%" stop-color="{GOLD_HI}"/><stop offset="100%" stop-color="#8C6E2E"/>
  </linearGradient>
  <filter id="glow" x="-60%" y="-60%" width="220%" height="220%"><feGaussianBlur stdDeviation="6"/></filter>
</defs>
<rect width="1280" height="720" fill="url(#bg)"/>
<text x="64" y="54" font-family="Inter" font-size="14" font-weight="600" letter-spacing="3" fill="{GOLD}">THE AI-NATIVE ORGANISATION</text>''')

    # connectors (behind)
    for nx, ny in nodes.values():
        parts.append(f'<line x1="{nx}" y1="{ny}" x2="{HUBX}" y2="{HUBY}" stroke="{GOLD}" stroke-opacity="0.20" stroke-width="1.25"/>')
    # faint system-boundary ring
    parts.append(f'<circle cx="{HUBX}" cy="{HUBY}" r="206" fill="none" stroke="{GOLD}" stroke-opacity="0.10" stroke-width="1" stroke-dasharray="2 7"/>')

    # hub
    parts.append(f'<circle cx="{HUBX}" cy="{HUBY}" r="{HUBR+3}" fill="none" stroke="{GOLD}" stroke-width="5" opacity="0.22" filter="url(#glow)"/>')
    parts.append(f'<circle cx="{HUBX}" cy="{HUBY}" r="{HUBR}" fill="url(#hub)" stroke="url(#goldring)" stroke-width="2.2"/>')
    parts.append(f'<text x="{HUBX}" y="258" text-anchor="middle" font-family="Inter" font-size="44" font-weight="800" fill="{WARM}" letter-spacing="1">AI</text>')
    parts.append(f'<text x="{HUBX}" y="286" text-anchor="middle" font-family="Inter" font-size="15" font-weight="600" letter-spacing="6" fill="{GOLD}">SYSTEMS</text>')

    pills = [("Skills &amp; Context", GOLD), ("Models &amp; Tools", GOLD), ("Security &amp; Governance", RED)]
    pw, ph = 230, 40
    for i, (label, accent) in enumerate(pills):
        py = 302 + i * 50
        cy = py + ph / 2
        fill = "#1E1714" if accent == RED else "#1B1B1B"
        parts.append(f'<rect x="{HUBX-pw/2}" y="{py}" width="{pw}" height="{ph}" rx="9" fill="{fill}" stroke="{accent}" stroke-opacity="{0.85 if accent==RED else 0.45}" stroke-width="1"/>')
        parts.append(f'<rect x="{HUBX-pw/2+10}" y="{cy-10}" width="3.2" height="20" rx="1.5" fill="{accent}"/>')
        parts.append(f'<text x="{HUBX+8}" y="{cy+5}" text-anchor="middle" font-family="Inter" font-size="14.5" font-weight="500" letter-spacing="0.3" fill="#EFEADF">{label}</text>')

    # nodes
    parts.append(node(*nodes["lead"], "compass"))
    parts.append(node(*nodes["del"], "heart"))
    parts.append(node(*nodes["grow"], "trending"))
    parts.append(node(*nodes["ops"], "gear"))

    def title(x, y, t, anchor):
        return f'<text x="{x}" y="{y}" text-anchor="{anchor}" font-family="Inter" font-size="19" font-weight="700" letter-spacing="2.5" fill="{TITLE}">{t}</text>'

    def desc(x, y, lines, anchor):
        return f'<text x="{x}" y="{y}" text-anchor="{anchor}" font-family="Inter" font-size="13.5" font-weight="400" letter-spacing="0.2" fill="{BODY}">{tspans(lines, x)}</text>'

    # Leadership (top, text right)
    parts.append(title(700, 100, "LEADERSHIP", "start"))
    parts.append(desc(700, 126, ["Sets direction, aligns strategy, and empowers the organisation."], "start"))
    # Operations (bottom, text right)
    parts.append(title(700, 622, "OPERATIONS", "start"))
    parts.append(desc(700, 648, ["Executes efficiently, optimises processes, and scales effectively."], "start"))
    # Delight (left, text right-aligned)
    parts.append(title(232, 356, "DELIGHT", "end"))
    parts.append(desc(232, 382, ["Understands needs, builds", "meaningful experiences,", "and earns loyalty."], "end"))
    # Growth (right, text left-aligned)
    parts.append(title(1048, 356, "GROWTH", "start"))
    parts.append(desc(1048, 382, ["Drives innovation, expands", "opportunities, and", "creates value."], "start"))

    # credits
    parts.append(f'<text x="64" y="694" font-family="Inter" font-size="11" font-weight="400" fill="{FAINT}">Framework adapted from Daniel Priestley</text>')
    parts.append(f'<text x="1216" y="694" text-anchor="end" font-family="Inter" font-size="11" font-weight="700" letter-spacing="1" fill="{FAINT}">BSTC</text>')

    parts.append('</svg>')
    svg = "\n".join(parts)
    with open(BASE + ".svg", "w") as f:
        f.write(svg)
    html = f'<!doctype html><html><head><meta charset="utf-8"><style>html,body{{margin:0;padding:0;background:#0B0B0B}}svg{{display:block}}</style></head><body>{svg}</body></html>'
    with open(BASE + ".html", "w") as f:
        f.write(html)
    print("Wrote", BASE + ".svg", "and .html")


def build_pptx():
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.dml.color import RGBColor
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0x0B, 0x0B, 0x0B)
    s.shapes.add_picture(BASE + ".png", 0, 0, width=Inches(13.333), height=Inches(7.5))
    prs.save(BASE + ".pptx")
    print("Wrote", BASE + ".pptx")


if __name__ == "__main__":
    mode = sys.argv[1] if len(sys.argv) > 1 else "svg"
    (build_pptx if mode == "pptx" else build_svg)()
